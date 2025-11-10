"""
Attachment Content Extraction Service

Handles extraction of text and structured content from various file types.
Supports PDF, Word, Excel, PowerPoint, CSV, TXT, and image files with OCR.
"""

import os
import json
import csv
from pathlib import Path
from typing import Dict, List, Optional, Any
import logging

# PDF extraction
try:
    import fitz  # PyMuPDF
    from pdfminer.high_level import extract_text as pdfminer_extract_text
    from pdfminer.pdfinterp import PDFResourceManager
    from pdfminer.pdfpage import PDFPage
    from pdfminer.converter import TextConverter
    from pdfminer.layout import LAParams
    from io import StringIO
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# Word document extraction
try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# Excel extraction
try:
    import openpyxl
    import pandas as pd
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

# PowerPoint extraction
try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

# OCR for images
try:
    import pytesseract
    from PIL import Image
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

# Setup logging
logger = logging.getLogger(__name__)


class AttachmentContentExtractor:
    """
    Service for extracting content from email attachments.
    
    Supports multiple file formats:
    - PDF: pdfminer.six, PyMuPDF
    - Word: python-docx
    - Excel: openpyxl, pandas
    - PowerPoint: python-pptx
    - CSV: pandas
    - TXT: built-in
    - Images: pytesseract OCR
    """
    
    def __init__(self):
        self.max_content_size = 100_000  # 100KB max content per file
        self.max_files = 20  # Max 20 attachments per email
        
        # File type handlers mapping
        self.handlers = {
            '.pdf': self._extract_pdf_content,
            '.doc': self._extract_doc_content,
            '.docx': self._extract_docx_content,
            '.txt': self._extract_txt_content,
            '.csv': self._extract_csv_content,
            '.xls': self._extract_excel_content,
            '.xlsx': self._extract_excel_content,
            '.ppt': self._extract_ppt_content,
            '.pptx': self._extract_pptx_content,
            '.jpg': self._extract_image_content,
            '.jpeg': self._extract_image_content,
            '.png': self._extract_image_content,
            '.gif': self._skip_file,  # Skip for now
            '.zip': self._skip_file,  # Skip for now
            '.rar': self._skip_file,  # Skip for now
        }
    
    def extract_attachments_content(self, attachment_files: List[Dict[str, str]]) -> List[Dict[str, Any]]:
        """
        Extract content from multiple attachment files.
        
        Args:
            attachment_files: List of dicts with 'file_name' and 'file_path'
            
        Returns:
            List of dicts with 'nome_file' and 'contenuto'
        """
        if not attachment_files:
            return []
        
        if len(attachment_files) > self.max_files:
            logger.warning(f"Too many attachments ({len(attachment_files)}), processing first {self.max_files}")
            attachment_files = attachment_files[:self.max_files]
        
        extracted_contents = []
        
        for file_info in attachment_files:
            try:
                file_name = file_info.get('file_name', 'unknown')
                file_path = file_info.get('file_path', '')
                
                print(f"📎 [EXTRACTOR] Processing attachment: {file_name}")
                
                if not file_path or not os.path.exists(file_path):
                    print(f"❌ [EXTRACTOR] File not found: {file_path}")
                    continue
                
                # Extract file extension
                file_ext = Path(file_name).suffix.lower()
                
                # Get appropriate handler
                handler = self.handlers.get(file_ext, self._unsupported_file)
                
                # Extract content
                content = handler(file_path, file_name)
                
                if content:
                    # Truncate content if too large
                    is_truncated = False
                    if len(str(content)) > self.max_content_size:
                        content = str(content)[:self.max_content_size] + "... [TRUNCATED]"
                        is_truncated = True
                    
                    extracted_contents.append({
                        "nome_file": file_name,
                        "contenuto": content
                    })
                    
                    print(f"✅ [EXTRACTOR] Successfully extracted content from {file_name} ({len(str(content))} chars)")
                    
                    # CONSOLE OUTPUT: Print extracted content preview
                    print(f"📄 [CONTENT PREVIEW] File: {file_name}")
                    print(f"   📋 Content Type: {file_ext}")
                    print(f"   📊 Content Length: {len(str(content))} characters")
                    print(f"   🔄 Truncated: {'Yes' if is_truncated else 'No'}")
                    print(f"   📝 Content Preview (first 500 chars):")
                    print(f"   " + "="*60)
                    preview_content = str(content)[:500]
                    for line in preview_content.split('\n'):
                        print(f"   {line}")
                    if len(str(content)) > 500:
                        print(f"   ... [showing first 500 of {len(str(content))} total characters]")
                    print(f"   " + "="*60)
                    print()
                else:
                    print(f"⚠️ [EXTRACTOR] No content extracted from {file_name}")
                
            except Exception as e:
                print(f"❌ [EXTRACTOR] Error processing {file_info.get('file_name', 'unknown')}: {e}")
                continue
        
        print(f"📎 [EXTRACTOR] Completed processing {len(extracted_contents)}/{len(attachment_files)} attachments")
        return extracted_contents
    
    def _extract_pdf_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract text content from PDF files."""
        try:
            # Try PyMuPDF first (faster)
            if PDF_AVAILABLE:
                try:
                    doc = fitz.open(file_path)
                    text = ""
                    for page in doc:
                        text += page.get_text()
                    doc.close()
                    
                    if text.strip():
                        return text.strip()
                except Exception as e:
                    print(f"⚠️ [EXTRACTOR] PyMuPDF failed for {file_name}, trying pdfminer: {e}")
            
                # Fallback to pdfminer
                try:
                    text = pdfminer_extract_text(file_path)
                    if text.strip():
                        return text.strip()
                except Exception as e:
                    print(f"❌ [EXTRACTOR] pdfminer also failed for {file_name}: {e}")
            
            return None
            
        except Exception as e:
            print(f"❌ [EXTRACTOR] PDF extraction error for {file_name}: {e}")
            return None
    
    def _extract_doc_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract text content from old Word documents (.doc)."""
        # .DOC files require more complex handling - for now return None
        # Could use python-docx2txt or antiword
        print(f"⚠️ [EXTRACTOR] .DOC files not yet supported: {file_name}")
        return None
    
    def _extract_docx_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract text content from Word documents (.docx)."""
        try:
            if not DOCX_AVAILABLE:
                return None
                
            doc = DocxDocument(file_path)
            
            # Extract text from paragraphs
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_parts.append(" | ".join(row_text))
            
            content = "\n".join(text_parts)
            return content if content.strip() else None
            
        except Exception as e:
            print(f"❌ [EXTRACTOR] DOCX extraction error for {file_name}: {e}")
            return None
    
    def _extract_txt_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract content from text files."""
        try:
            # Try different encodings
            encodings = ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    if content.strip():
                        return content.strip()
                except UnicodeDecodeError:
                    continue
            
            print(f"❌ [EXTRACTOR] Could not decode text file {file_name}")
            return None
            
        except Exception as e:
            print(f"❌ [EXTRACTOR] TXT extraction error for {file_name}: {e}")
            return None
    
    def _extract_csv_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract structured content from CSV files."""
        try:
            if not EXCEL_AVAILABLE:
                # Fallback to basic CSV parsing
                return self._extract_csv_basic(file_path, file_name)
            
            # Use pandas for better CSV handling
            try:
                df = pd.read_csv(file_path, encoding='utf-8')
            except UnicodeDecodeError:
                df = pd.read_csv(file_path, encoding='latin-1')
            
            # Convert to structured text format
            if df.empty:
                return None
            
            # Limit rows for large CSV files
            if len(df) > 100:
                df = df.head(100)
                truncated = True
            else:
                truncated = False
            
            # Create structured text representation
            content_parts = []
            content_parts.append(f"CSV File: {file_name}")
            content_parts.append(f"Columns: {', '.join(df.columns.tolist())}")
            content_parts.append(f"Rows: {len(df)}")
            content_parts.append("")
            
            # Add sample data
            content_parts.append("Sample Data:")
            for idx, row in df.iterrows():
                row_text = " | ".join([f"{col}: {val}" for col, val in row.items() if pd.notna(val)])
                content_parts.append(row_text)
            
            if truncated:
                content_parts.append("... [TRUNCATED - showing first 100 rows]")
            
            return "\n".join(content_parts)
            
        except Exception as e:
            print(f"❌ [EXTRACTOR] CSV extraction error for {file_name}: {e}")
            return None
    
    def _extract_csv_basic(self, file_path: str, file_name: str) -> Optional[str]:
        """Basic CSV extraction without pandas."""
        try:
            content_parts = []
            
            with open(file_path, 'r', encoding='utf-8', newline='') as f:
                reader = csv.reader(f)
                rows = list(reader)
            
            if not rows:
                return None
            
            # Header
            content_parts.append(f"CSV File: {file_name}")
            if rows:
                content_parts.append(f"Columns: {', '.join(rows[0])}")
                content_parts.append(f"Rows: {len(rows) - 1}")
                content_parts.append("")
            
            # Sample data (first 50 rows)
            sample_rows = rows[:51]  # Header + 50 data rows
            for row in sample_rows:
                content_parts.append(" | ".join(row))
            
            if len(rows) > 51:
                content_parts.append("... [TRUNCATED]")
            
            return "\n".join(content_parts)
            
        except Exception as e:
            print(f"❌ [EXTRACTOR] Basic CSV extraction error for {file_name}: {e}")
            return None
    
    def _extract_excel_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract structured content from Excel files."""
        try:
            if not EXCEL_AVAILABLE:
                return None
            
            # Read Excel file
            try:
                xl_file = pd.ExcelFile(file_path)
                sheet_names = xl_file.sheet_names
            except Exception as e:
                print(f"❌ [EXTRACTOR] Cannot read Excel file {file_name}: {e}")
                return None
            
            content_parts = []
            content_parts.append(f"Excel File: {file_name}")
            content_parts.append(f"Sheets: {', '.join(sheet_names)}")
            content_parts.append("")
            
            # Process each sheet (limit to first 3 sheets)
            for sheet_name in sheet_names[:3]:
                try:
                    df = pd.read_excel(file_path, sheet_name=sheet_name)
                    
                    if df.empty:
                        continue
                    
                    content_parts.append(f"Sheet: {sheet_name}")
                    content_parts.append(f"Columns: {', '.join(df.columns.tolist())}")
                    content_parts.append(f"Rows: {len(df)}")
                    
                    # Add sample data (first 20 rows)
                    sample_df = df.head(20)
                    for idx, row in sample_df.iterrows():
                        row_text = " | ".join([f"{col}: {val}" for col, val in row.items() if pd.notna(val)])
                        content_parts.append(row_text)
                    
                    if len(df) > 20:
                        content_parts.append("... [TRUNCATED - showing first 20 rows]")
                    
                    content_parts.append("")
                    
                except Exception as e:
                    content_parts.append(f"Error reading sheet {sheet_name}: {e}")
                    continue
            
            return "\n".join(content_parts) if content_parts else None
            
        except Exception as e:
            print(f"❌ [EXTRACTOR] Excel extraction error for {file_name}: {e}")
            return None
    
    def _extract_ppt_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract text content from old PowerPoint files (.ppt)."""
        # .PPT files require more complex handling - for now return None
        print(f"⚠️ [EXTRACTOR] .PPT files not yet supported: {file_name}")
        return None
    
    def _extract_pptx_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract text content from PowerPoint files (.pptx)."""
        try:
            if not PPTX_AVAILABLE:
                return None
            
            prs = Presentation(file_path)
            
            content_parts = []
            content_parts.append(f"PowerPoint File: {file_name}")
            content_parts.append(f"Slides: {len(prs.slides)}")
            content_parts.append("")
            
            # Extract text from each slide
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = []
                
                # Extract text from shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                
                if slide_text:
                    content_parts.append(f"Slide {slide_num}:")
                    content_parts.extend(slide_text)
                    content_parts.append("")
            
            content = "\n".join(content_parts)
            return content if content.strip() else None
            
        except Exception as e:
            print(f"❌ [EXTRACTOR] PPTX extraction error for {file_name}: {e}")
            return None
    
    def _extract_image_content(self, file_path: str, file_name: str) -> Optional[str]:
        """Extract text content from images using OCR."""
        try:
            if not OCR_AVAILABLE:
                print(f"⚠️ [EXTRACTOR] OCR not available for {file_name}")
                return None
            
            # Open image
            image = Image.open(file_path)
            
            # Extract text using OCR
            text = pytesseract.image_to_string(image)
            
            if text.strip():
                content_parts = []
                content_parts.append(f"Image File: {file_name}")
                content_parts.append(f"OCR Extracted Text:")
                content_parts.append(text.strip())
                return "\n".join(content_parts)
            
            return None
            
        except Exception as e:
            print(f"❌ [EXTRACTOR] OCR extraction error for {file_name}: {e}")
            return None
    
    def _skip_file(self, file_path: str, file_name: str) -> Optional[str]:
        """Skip file processing (for unsupported types like ZIP, RAR, GIF)."""
        print(f"⚠️ [EXTRACTOR] Skipping file type: {file_name}")
        return None
    
    def _unsupported_file(self, file_path: str, file_name: str) -> Optional[str]:
        """Handle unsupported file types."""
        print(f"⚠️ [EXTRACTOR] Unsupported file type: {file_name}")
        return None


# Create singleton instance
attachment_extractor = AttachmentContentExtractor()